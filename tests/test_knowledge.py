"""知识库：分类、上传(txt/md/csv/docx)、分块建索引、TF-IDF 检索、删除。"""
import io

from conftest import auth


def _upload(client, token, kb_id, filename, content: bytes):
    return client.post(f"/api/kb/{kb_id}/documents", headers=auth(token),
                       files={"file": (filename, io.BytesIO(content), "application/octet-stream")})


def test_kb_categories_and_seed(client, emp_token):
    cats = client.get("/api/kb/categories", headers=auth(emp_token)).json()
    assert cats == ["产品资料", "公司SOP", "供应商资料", "运营经验", "质量标准"]
    kbs = client.get("/api/kb", headers=auth(emp_token)).json()
    assert len(kbs) >= 2
    assert any(k["category"] == "公司SOP" and k["doc_count"] >= 1 for k in kbs)


def test_create_kb_validation(client, admin_token, emp_token):
    assert client.post("/api/kb", headers=auth(emp_token),
                       json={"name": "x", "category": "产品资料"}).status_code == 403
    assert client.post("/api/kb", headers=auth(admin_token),
                       json={"name": "x", "category": "不存在的分类"}).status_code == 400
    r = client.post("/api/kb", headers=auth(admin_token),
                    json={"name": "供应商资料库", "category": "供应商资料", "description": "供应商档案"})
    assert r.status_code == 200


def test_upload_txt_md_csv_and_search(client, admin_token, emp_token):
    kb_id = client.post("/api/kb", headers=auth(admin_token),
                        json={"name": "运营经验库", "category": "运营经验"}).json()["id"]

    r = _upload(client, admin_token, kb_id, "广告优化经验.md",
                "# 亚马逊PPC广告优化\n\n当ACOS高于30%时，应检查关键词竞价并否定无效搜索词。".encode("utf-8"))
    assert r.status_code == 200 and r.json()["chunk_count"] >= 1

    _upload(client, admin_token, kb_id, "备注.txt",
            "站外推广优先选择Facebook群组与红人营销，转化率通常高于展示广告。".encode("utf-8"))
    _upload(client, admin_token, kb_id, "关键词.csv",
            "关键词,搜索量,转化率\n保温杯,52000,8.5%\n便携榨汁杯,31000,7.2%\n".encode("utf-8"))

    docs = client.get(f"/api/kb/{kb_id}/documents", headers=auth(emp_token)).json()
    assert len(docs) == 3
    assert sum(d["chunk_count"] for d in docs) >= 3

    # 员工也能检索（AI回答依赖它）
    r = client.get(f"/api/kb/{kb_id}/search", headers=auth(emp_token),
                   params={"q": "ACOS高于30%应该怎么优化广告", "top_k": 3})
    assert r.status_code == 200
    results = r.json()["results"]
    assert results and results[0]["filename"] == "广告优化经验.md"
    assert results[0]["score"] > 0
    assert "否定无效搜索词" in results[0]["text"]

    # CSV 内容也可检索
    results = client.get(f"/api/kb/{kb_id}/search", headers=auth(emp_token),
                         params={"q": "保温杯 搜索量"}).json()["results"]
    assert any(r["filename"] == "关键词.csv" for r in results)


def test_upload_docx(client, admin_token, emp_token):
    from docx import Document

    doc = Document()
    doc.add_heading("质量检验标准", level=1)
    doc.add_paragraph("保温杯出厂前必须做24小时密封性测试，漏水率不得超过千分之三。")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "外观划痕"
    table.rows[0].cells[1].text = "AQL 2.5"
    buf = io.BytesIO()
    doc.save(buf)

    kb_id = client.post("/api/kb", headers=auth(admin_token),
                        json={"name": "质量标准库", "category": "质量标准"}).json()["id"]
    r = _upload(client, admin_token, kb_id, "质检标准.docx", buf.getvalue())
    assert r.status_code == 200 and r.json()["chunk_count"] >= 1

    results = client.get(f"/api/kb/{kb_id}/search", headers=auth(emp_token),
                         params={"q": "保温杯密封性测试漏水率要求"}).json()["results"]
    assert results and "千分之三" in results[0]["text"]
    # 表格内容也进了索引
    all_text = " ".join(x["text"] for x in client.get(f"/api/kb/{kb_id}/search", headers=auth(emp_token),
                                                       params={"q": "外观划痕 AQL"}).json()["results"])
    assert "AQL" in all_text


def test_upload_unsupported_and_empty(client, admin_token):
    kb_id = client.get("/api/kb", headers=auth(admin_token)).json()[0]["id"]
    assert _upload(client, admin_token, kb_id, "恶意.exe", b"MZ...").status_code == 400
    assert _upload(client, admin_token, kb_id, "空.txt", b"").status_code == 400
    assert _upload(client, admin_token, 99999, "a.txt", b"hello world hello world").status_code == 404


def test_delete_document_and_kb(client, admin_token):
    kb_id = client.post("/api/kb", headers=auth(admin_token),
                        json={"name": "临时库", "category": "产品资料"}).json()["id"]
    doc = _upload(client, admin_token, kb_id, "临时.txt", "临时内容ABC123".encode("utf-8")).json()
    assert client.delete(f"/api/kb/{kb_id}/documents/{doc['doc_id']}", headers=auth(admin_token)).status_code == 200
    assert client.get(f"/api/kb/{kb_id}/search", headers=auth(admin_token),
                      params={"q": "临时内容ABC123"}).json()["results"] == []
    assert client.delete(f"/api/kb/{kb_id}", headers=auth(admin_token)).status_code == 200

    # 被助手绑定的知识库不允许删除
    kbs = client.get("/api/kb", headers=auth(admin_token)).json()
    sop = next(k for k in kbs if k["name"] == "公司SOP库")
    r = client.delete(f"/api/kb/{sop['id']}", headers=auth(admin_token))
    assert r.status_code == 400
